#!/usr/bin/env python3
"""Update the retained sociogram report from the published graph and official catalogue."""
from pathlib import Path
from collections import Counter
import json, re, copy, hashlib, subprocess, sys
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
PROJECT=ROOT.parents[1]
SRC=ROOT/'descarregues/Sociograma_comunitari_CAP_Salt1_Salt2_2026.docx'
if not SRC.exists():
    SRC=PROJECT/'reference/sociograma_salt_2026-10-08/sociograma_salt_2026-10-08/descarregues/Sociograma_comunitari_CAP_Salt1_Salt2_2026.docx'
OUT=ROOT/'descarregues/Sociograma_Salt_Observatori_complet_2026-10-17.docx'
TMP=ROOT/'tmp/documents'
TMP.mkdir(parents=True,exist_ok=True)
cat=json.loads((ROOT/'data/observatori-cataleg.json').read_text())
assert len(cat['fitxes']) == 87, 'The report requires the complete official catalogue'
actors=json.loads((ROOT/'data/actors.json').read_text())
rels=json.loads((ROOT/'data/relacions.json').read_text())
high=json.loads((ROOT/'data/observatori.json').read_text())
d=Document(SRC)
source_paragraphs=list(d.paragraphs)
assert len(source_paragraphs)==242, 'Reconcile paragraph anchors if the retained report changes'
# Stable semantic anchors against the retained 7 October report (the GitHub blob).
# Keys retain the editorial slot numbering used in the preceding 6 October draft.
anchor_map={i:i for i in range(6)}
anchor_map.update({i:i+1 for i in range(6,21)})
anchor_map.update({i:i+2 for i in range(21,66)})
anchor_map.update({i:i+8 for i in range(66,70)})
anchor_map[70]=79
anchor_map.update({i:i+9 for i in range(71,77)})
anchor_map.update({77:86,78:87,79:88,80:89,81:90,82:91,83:92,84:94})
anchor_map.update({i:i+10 for i in range(85,91)})
anchor_map.update({i:i+11 for i in range(91,148)})
anchor_map.update({i:i+23 for i in range(148,219)})
original={i:source_paragraphs[j] for i,j in anchor_map.items()}
original_tables=list(d.tables)
source_hash=hashlib.sha256(SRC.read_bytes()).hexdigest()
PAGE_WIDTH=17.59

def txt(p,text):
    p.text=text
    return p

def para(text='',style=None):
    p=d.add_paragraph(text,style)
    p.paragraph_format.space_after=Pt(5)
    p.paragraph_format.line_spacing=1.08
    return p

def heading(text,level=1):
    p=d.add_heading(text,level)
    p.paragraph_format.space_before=Pt(12 if level==1 else 8)
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.keep_with_next=True
    return p

def link(p,text,url,size=None):
    h=OxmlElement('w:hyperlink')
    h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color')
    color.set(qn('w:val'),'174F76');props.append(color)
    if size:
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(round(size*2)));props.append(sz)
    r.append(props);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def table(headers,rows,widths=None,size=9):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i,label in enumerate(headers):t.rows[0].cells[i].text=label
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):cells[i].text=str(value)
    format_table(t,widths,size)
    para('')
    return t

def format_table(t,widths=None,size=9):
    if widths:
        for i,w in enumerate(widths):
            t.columns[i].width=Cm(w)
            for cell in t.columns[i].cells:cell.width=Cm(w)
    pr=t._tbl.tblPr
    borders=pr.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
    for tag in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        for fixed in list(trpr.findall(qn('w:trHeight'))):trpr.remove(fixed)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        for ci,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            for old in list(cp.findall(qn('w:shd'))):cp.remove(old)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'244C66' if ri==0 else ('F4F6F7' if ri%2 else 'FFFFFF'));cp.append(sh)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1)
                p.paragraph_format.space_before=Pt(1)
                p.paragraph_format.line_spacing=1.0
                p.paragraph_format.keep_with_next=False
                for r in p.runs:
                    r.font.size=Pt(size);r.font.bold=(ri==0);r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')

def replace_table(old,headers,rows,widths,size=8):
    new=table(headers,rows,widths,size)
    old._tbl.addprevious(new._tbl);old._tbl.getparent().remove(old._tbl)
    return new

def num(s):
    s=str(s).strip().replace('\xa0','').replace(' ','').replace('%','')
    if s in ('','-','..','...','—','n/d'):return None
    try:return float(s.replace('.','').replace(',','.')) if ',' in s else float(s)
    except ValueError:return None

def fmt(x,dec=0):
    return f'{x:,.{dec}f}'.replace(',','_').replace('.',',').replace('_','.')

def find_fiche(fragment):
    fs=[x for x in cat['fitxes'] if x['url'].rstrip('/').endswith('/'+fragment)]
    if not fs:fs=[x for x in cat['fitxes'] if fragment.lower() in (x['titol']+' '+x['url']).lower()]
    if not fs and fragment=='llars-per-nombre':
        fs=[x for x in cat['fitxes'] if any(w in x['titol'].lower() for w in ['llars','domicilis']) and any(w in x['titol'].lower() for w in ['grandària','nombre de persones','nombre de membres','ocupants']) and 'unipersonals' not in x['titol'].lower()]
    return fs[0] if fs else None

def choose_variant(f):
    vs=f.get('variants',[])
    default=next((v for v in vs if v.get('hash')==f.get('variant_per_defecte') and v.get('tables')),None)
    if default:return default
    municipal=[v for v in vs if v.get('tables') and (v.get('ambit','').lower()=='municipi' or '/Municipi/' in v.get('hash','') or 'municipi/Municipi' in v.get('hash',''))]
    neutral=[v for v in municipal if str(v.get('variant') or '').casefold() in ('','total','general','tota la població','tots','població total')]
    if neutral:return neutral[0]
    for v in vs:
        if v.get('tables') and (v.get('ambit','').lower()=='municipi' or '/Municipi/' in v.get('hash','') or 'municipi/Municipi' in v.get('hash','')):
            return v
    return next((v for v in vs if v.get('tables')),None)

def table_cols(t):
    cols=t.get('columnes')
    if cols:return cols
    headers=t.get('headers',[])
    if not headers:return []
    if isinstance(headers[0],str):return headers
    n=max(map(len,headers));return [' · '.join(str(r[i]) for r in headers if i<len(r) and r[i]) for i in range(n)]

def latest_summary(f,maxrows=5):
    v=choose_variant(f)
    if not v:return None,None,[]
    t=v['tables'][0];cols=table_cols(t)
    years=[int(y) for c in cols for y in re.findall(r'\b(?:19|20)\d{2}\b',str(c))]
    year=max(years) if years else None
    idx=[i for i,c in enumerate(cols) if year and str(year) in str(c)]
    if not idx:idx=list(range(1,min(len(cols),3)))
    rows=[]
    for r in t.get('rows',[])[:maxrows]:
        if not r:continue
        rows.append([r[0]]+[r[i] if i<len(r) else '' for i in idx])
    return year,[cols[0] if cols else 'Grup']+[cols[i] for i in idx],rows

def source_line(f):
    p=para('Font: '+(str(f.get('font','')).strip() or 'Observatori de Salt')+'. ')
    for r in p.runs:r.font.size=Pt(8.5)
    link(p,'Fitxa oficial',f['url'])
    return p

def chart_png(drawing,name):
    from reportlab.graphics import renderPDF
    pdf=TMP/(name+'.pdf');png=TMP/(name+'.png')
    renderPDF.drawToFile(drawing,str(pdf))
    bundled=Path(sys.executable).parents[2]/'bin/override/pdftoppm'
    subprocess.run([str(bundled),'-png','-singlefile','-r','150',str(pdf),str(png.with_suffix(''))],check=True,capture_output=True)
    return png

def population_pyramid(f):
    """Only draw when sex labels, age rows and the population total reconcile."""
    series={}
    target=2025
    neutral=choose_variant(f)
    if neutral:
        for t in neutral.get('tables',[]):
            cols=table_cols(t)
            for sex in ['Dones','Homes']:
                choices=[i for i,c in enumerate(cols) if str(target) in str(c) and sex.casefold() in str(c).casefold()]
                if choices:
                    ci=choices[0];points={}
                    for r in t.get('rows',[]):
                        if not r or ci>=len(r):continue
                        label=str(r[0]);value=num(r[ci])
                        if re.search(r'\d',label) and value is not None:points[label]=value
                    if len(points)>=10:series[sex]=points
    for v in ([] if len(series)==2 else f.get('variants',[])):
        identity=' '.join(str(v.get(k,'')) for k in ['variant','hash','url']).casefold()
        sex='Dones' if 'dones' in identity else ('Homes' if 'homes' in identity else None)
        if not sex or ('Municipi' not in v.get('hash','') and v.get('ambit','').lower()!='municipi'):
            continue
        for t in v.get('tables',[]):
            cols=table_cols(t)
            choices=[i for i,c in enumerate(cols) if str(target) in str(c) and '%' not in str(c)]
            if not choices:continue
            ci=choices[0];points={}
            for r in t.get('rows',[]):
                if not r or ci>=len(r):continue
                label=str(r[0]);value=num(r[ci])
                if re.search(r'\d',label) and value is not None:points[label]=value
            if len(points)>=10:series[sex]=points
    if set(series)!={'Homes','Dones'} or set(series['Homes'])!=set(series['Dones']):return False
    labels=list(series['Homes'])
    total=sum(series[s][a] for s in series for a in labels)
    if abs(total-35106)>1:return False
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.lib.colors import HexColor
    labels.sort(key=lambda a:int(re.search(r'\d+',a).group()))
    drawing=Drawing(660,410);chart=HorizontalBarChart()
    chart.x=120;chart.y=45;chart.width=520;chart.height=340
    chart.data=[[-series['Homes'][a] for a in labels],[series['Dones'][a] for a in labels]]
    chart.categoryAxis.categoryNames=labels
    chart.categoryAxis.style='stacked';chart.categoryAxis.labels.fontSize=8
    chart.categoryAxis.labelAxisMode='low'
    chart.categoryAxis.labels.boxAnchor='e';chart.categoryAxis.labels.dx=-5
    chart.valueAxis.labelTextFormat=lambda x:fmt(abs(x))
    chart.valueAxis.labels.fontSize=9
    bound=max(max(series[s].values()) for s in series)*1.12
    chart.valueAxis.valueMin=-bound;chart.valueAxis.valueMax=bound
    chart.bars[0].fillColor=HexColor('#246A7A');chart.bars[1].fillColor=HexColor('#BB6752')
    chart.bars.strokeWidth=0;chart.barSpacing=0;chart.groupSpacing=1
    drawing.add(chart);drawing.add(String(240,395,'Homes',fontSize=11,fillColor=HexColor('#246A7A')))
    drawing.add(String(470,395,'Dones',fontSize=11,fillColor=HexColor('#BB6752')))
    drawing.add(String(280,10,'Persones empadronades',fontSize=10))
    png=chart_png(drawing,'population_pyramid');d.add_picture(str(png),width=Cm(16.8))
    para('Piràmide municipal de 2025. Les dues sèries sumen 35.106 habitants; sexe i grups d’edat segons la classificació de la font.')
    return True

# Correct superseded numbers and retain the original scope and narrative.
txt(original[0],'Sociograma comunitari dels CAP Salt 1 i Salt 2')
original[0].style=d.styles['Title']
original[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in original[0].runs:r.font.size=Pt(24);r.font.bold=True;r.font.color.rgb=RGBColor(0,0,0)
txt(original[2],"Actualització del 17 d'octubre de 2026. Dades municipals de l'Observatori de Salt i catàleg complet de fitxes. Les propostes i les relacions hipotètiques requereixen contrast amb l'EAP i els agents comunitaris.")
txt(original[10],"Exclusió de municipis. L'ABS 199 inclou també Aiguaviva, Bescanó, Fornells de la Selva i Vilablareix. Aquests municipis queden fora del sociograma, que se centra en Salt. L'Observatori publica 35.106 habitants per a 2025. Aquest padró municipal no equival a població assignada als CAP ni es pot combinar directament amb un total d'ABS d'una explotació diferent. Els indicadors AQuAS d'ABS inclouen els altres municipis.")
txt(original[12],"Proximitat als CAP. La zonificació geomètrica i el CAP més proper són una orientació espacial, no una assignació assistencial oficial. A la base actual, 87 actors amb coordenades queden més a prop de Salt 1 i 33 de Salt 2. La direcció de l'EAP ha de validar el repartiment real de carrers, UBA i població assignada.")
txt(original[13],"Barris. L'Observatori municipal publica població i altres indicadors per barris, incloses sèries històriques. El catàleg incorporat permet consultar aquests desglossaments. Els barris de l'Observatori, els sectors dels Consells de Barri, les seccions censals i les zones dels CAP són unitats diferents; no s'assumeix cap correspondència. El camp zona/barri dels actors continua sent orientatiu quan no té una delimitació contrastada.")
txt(original[39],f"La base publicada conté {len(actors)} nodes, inclosos els 256 actors i els 3 agregadors de la vista general, i {len(rels)} relacions: 322 funcionals documentades, 100 estructurals documentades i 56 hipòtesis a validar. L'annex A i el resum per esferes s'han sincronitzat amb aquesta base. El registre municipal d'entitats continua sent una font complementària i no acredita per si sol activitat actual.")
txt(original[41],'Per forma: 19 triangles, 225 rectangles i 15 cercles. Per estat al dataset: 252 verificats, 5 probables i 2 a validar. Aquest estat correspon a les fonts recollides per cada actor i no substitueix la validació comunitària.')
txt(original[64],"Figura 4. Mapa d'actius de l'informe de base del 7 d'octubre. La base actual inclou 120 actors amb coordenades. Les zones són de proximitat i no representen l'assignació oficial dels CAP. Consulteu el visor per al mapa actualitzat.")
for n in [4,60,62]:txt(original[n],original[n].text+' Imatge conservada de l’informe de base del 7 d’octubre; consulteu el visor per a la base actualitzada.')
txt(original[67],"L'Ajuntament de Salt i la Direcció de l'EAP articulen respectivament bona part de la xarxa municipal i sanitària. El grau i els vincles de cada node s'han de consultar al visor actualitzat, que distingeix les relacions documentades de les hipòtesis.")
# Preserve the confirmed coordination structures and activities added on 7 October.
txt(original[77],original[77].text.replace('buit confirmat','relació pendent de contrast')+' L’absència d’una activitat a la memòria no acredita absència de col·laboració real.')
txt(original[111],"Propostes inicials de treball a revisar amb la Comissió de Salut Comunitària. La prioritat indicada és una proposta, no un acord vigent ni una conclusió de l'Observatori. Els actors marcats (H) representen relacions proposades o pendents de validar; cal contrastar-les amb el dataset actual.")
txt(original[136],'Les 56 relacions marcades com a hipòtesis en la base actual (annex B): confirmar, canviar el tipus o eliminar.')
txt(original[147],'Annex C Fonts del sociograma')
txt(original[218],"Dades i visor: https://cvragar.github.io/sociograma-salt/. Llicència dels indicadors AQuAS: CC BY-NC-ND 4.0. Les fonts del catàleg municipal s'identifiquen a l'annex D i als fitxers de dades descarregables.")

# Synchronize category summaries while retaining the original chapter structure.
level='polític'
for i in range(42,57):
    p=original[i]
    if p.style.name=='Heading 2':
        level='polític' if 'polític' in p.text else ('tècnic' if 'tècnic' in p.text else 'social')
    elif ': ' in p.text and '(' in p.text.split(': ')[0]:
        sphere=re.sub(r'\s*\(\d+\)$','',p.text.split(': ')[0])
        group=[a for a in actors if a['nivell']==level and a['esfera']==sphere]
        if group:txt(p,sphere+f' ({len(group)}): '+'; '.join(a['nom']+(f" [{a['estat_verificacio']}]" if a['estat_verificacio']!='verificat' else '') for a in group)+'.')

values={x['id']:x for x in high['indicadors']}
original_tables[0].cell(1,1).text="Salt: 35.106 habitants (Observatori, 2025); densitat bruta 5.311,7 hab./km². La temporalitat és anual i no se n'explicita el dia de tall. Població assignada a l'ABS: 42.298 (AQuAS 2022), àmbit i període diferents."
original_tables[0].cell(1,2).text='Observatori de Salt 2025; AQuAS 2022'
original_tables[0].cell(2,1).text="Observatori 2025: 12.899 persones amb nacionalitat estrangera (36,7%) i 16.298 nascudes a l'estranger (46,4%). País de naixement i nacionalitat són variables diferents. Aquestes dades no identifiquen llengües, creences, situació administrativa ni necessitats individuals."
original_tables[0].cell(2,2).text='Observatori de Salt 2025'
original_tables[0].cell(6,1).text="Observatori 2025: 4.800 persones de 65 anys o més (13,7%); 2.133 domicilis amb una sola persona empadronada, dels quals 360 al grup literal «Més de 80 anys». Viure sol no equival a soledat no desitjada."
original_tables[0].cell(6,2).text='Observatori de Salt 2025'
old=original_tables[0].cell(5,1).text
original_tables[0].cell(5,1).text=old.replace('la qual cosa apunta a una possible infrautilització de la salut mental especialitzada.','una diferència que cal contextualitzar abans d’atribuir-la a infrautilització.')
original_tables[3].cell(5,0).text='Delimitació i correspondència entre barris, seccions censals i cobertura dels CAP'
original_tables[3].cell(5,2).text="Cartografia i llista de carrers dels barris de l'Observatori, i correspondència amb seccions censals. El padró per barris ja és consultable al catàleg municipal."
original_tables[2].cell(1,3).text="La xarxa documental requereix validació de circuits i prioritats; contrastar la planificació local vigent."
original_tables[2].cell(2,3).text="Obesitat d'ABS 2022 i desigualtat territorial. El públic i el disseny s'han d'acordar amb participants i agents, sense inferir necessitats pel país d'origen."
original_tables[2].cell(5,3).text='Observatori 2025: 2.133 llars unipersonals. Cal explorar suport social i preferències, sense equiparar viure sol amb soledat.'
spheres=list(dict.fromkeys(a['esfera'] for a in actors))
rows=[]
for s in spheres:
    counts=[sum(a['esfera']==s and a['nivell']==lv for a in actors) for lv in ['polític','tècnic','social']]
    rows.append([s,*counts,sum(counts)])
rows.append(['TOTAL',*[sum(a['nivell']==lv for a in actors) for lv in ['polític','tècnic','social']],len(actors)])
replace_table(original_tables[1],['Esfera','Polític','Tècnic','Social','Total'],rows,[9.1,2.1,2.1,2.1,2.1],9)
actor_rows=[[a['id'],a['nom'],a['tipus_alberich'][:5]+'.',a['nivell'],a['esfera'],a['estat_verificacio'],str(a.get('cap_mes_proper') or 'Sense dada').replace('CAP ','')] for a in actors]
replace_table(original_tables[7],['ID','Actor','Forma','Nivell','Esfera','Estat','Proximitat'],actor_rows,[1.0,6.5,1.3,1.4,3.5,1.7,2.1],8)
hyp=[r for r in rels if r.get('estat')!='documentada']
hr=[[r['origen']+' '+r.get('origen_nom',''),r['desti']+' '+r.get('desti_nom',''),r['tipus_relacio'],r.get('evidencia','')] for r in hyp]
replace_table(original_tables[8],['Origen','Destí','Tipus','Comentari'],hr,[5.2,5.2,1.6,5.5],8)

# The chapter is authored at the end and moved after the cover metadata.
start=len(d.element.body)
p=heading('Actualització amb les dades de l’Observatori')
m=cat['metadades'];fiches=cat['fitxes']
withdata=sum(f.get('estat')=='disponible' for f in fiches)
para(f"L'actualització incorpora les {len(fiches)} fitxes del catàleg oficial consultat el 17 d'octubre de 2026. {withdata} fitxes contenen taules recuperades. L'annex D identifica les fitxes, la cobertura temporal i l'estat de les dades; les sèries i els desglossaments complets es consulten al visor i es descarreguen en JSON i CSV.")
para("L'objectiu és orientar la planificació comunitària dels CAP de Salt amb informació del municipi. El padró municipal no equival a població assignada a cada CAP, i els barris no són zones assistencials. Cada indicador conserva l'any i la font; la data d'aquesta revisió no converteix totes les dades en dades de 2026.")
heading('Població i prioritats de treball',2)
table(['Indicador municipal','Valor','Any'],[
['Població empadronada','35.106 habitants','2025'],['De 0 a 15 anys','6.455 persones · 18,4%','2025'],['De 16 a 64 anys','23.851 persones · 67,9%','2025'],['65 anys i més','4.800 persones · 13,7%','2025'],['Persones nascudes a l’estranger','16.298 · 46,4%','2025'],['Persones amb nacionalitat estrangera','12.899 · 36,7%','2025'],['Domicilis amb una sola persona empadronada','2.133','2025'],['Ocupació mitjana dels domicilis','3,32 persones per domicili','2025'],['Densitat bruta de població','5.311,7 habitants per km²','2025']], [9.2,5.7,2.6])
p=para('Fonts: ')
for j,key in enumerate(['population','foreign_born','single_households','household_occupancy','gross_density']):
    if j:p.add_run(' · ')
    link(p,values[key]['nom'],values[key]['url'])
para("La composició per edat ajuda a dimensionar activitats amb escoles, famílies i gent gran. La diversitat de naixements i nacionalitats fa pertinent acordar canals accessibles de participació; no permet deduir la llengua, la religió ni les necessitats de cada persona. Les propostes següents són decisions de treball a validar, no conclusions causals del portal.")

# Selected raw tables are placed without altering published values.
selected=[
('estructura-de-la-poblacio-per-grups-dedat-quinquennals-i-sexe','Estructura de la població per edat i sexe',22),
('llars-unipersonals-per-grups-dedat','Llars unipersonals segons l’edat',15),
('llars-per-nombre','Composició de les llars',12),
]
used=set()
for fragment,title,maxrows in selected:
    f=find_fiche(fragment)
    if not f or f['id'] in used:continue
    used.add(f['id'])
    year,cols,rows=latest_summary(f,maxrows)
    if not rows:continue
    heading(title,2)
    intro=para('Dades municipals'+(f' de {year}' if year else '')+'. Valors de la taula oficial.')
    intro.paragraph_format.keep_with_next=True
    pyramid=population_pyramid(f) if 'quinquennals-i-sexe' in fragment else False
    if not pyramid and cols and len(cols)<=5:
        if not cols[0]:cols[0]='Persones per domicili' if fragment=='llars-per-nombre' else 'Grup d’edat'
        if v:=choose_variant(f):
            if v['tables'][0].get('format_numeric')=='decimal-punt':
                for row in rows:
                    for ci in range(1,len(row)):
                        raw=str(row[ci]);n=num(raw)
                        if n is not None:
                            decimals=len(raw.rstrip('%').split('.')[1]) if '.' in raw else 0
                            row[ci]=fmt(n,decimals)+('%' if '%' in raw else '')
        widths=[6.2]+[(17.5-6.2)/(len(cols)-1)]*(len(cols)-1)
        table(cols,rows,widths,9)
    source_line(f)
    if 'unipersonals' in fragment:para("La categoria «Més de 80 anys» es manté literalment. El portal no concreta el límit del grup. Viure sol és una característica del domicili i no un diagnòstic de soledat no desitjada.")

heading('Sèries històriques i comparacions',2)
para("Les fitxes conserven tots els períodes i àmbits recuperats, incloses les variants municipals i per barris o seccions quan la font les publica. Per comparar anys cal mantenir la mateixa variable, unitat, àmbit i font. Un canvi de sèrie, de tall del padró o de definició pot explicar diferències entre portals; no es calcula cap tendència combinant explotacions diferents.")
f=find_fiche('estructura-de-la-poblacio-per-grans-grups-dedat')
if f:
    v=choose_variant(f)
    if v:
        t=v['tables'][0];cols=table_cols(t);total=next((r for r in t.get('rows',[]) if r and str(r[0]).lower() in ('total','població total','poblacio total')),None)
        vals=[]
        if total:
            for i,c in enumerate(cols):
                yy=re.findall(r'\b(?:19|20)\d{2}\b',str(c))
                if yy and '%' not in c and i<len(total):
                    n=num(total[i])
                    if n is not None:vals.append((int(yy[0]),n))
        vals=sorted(set(vals))
        if len(vals)>1:
            from reportlab.graphics.shapes import Drawing, String
            from reportlab.graphics.charts.lineplots import LinePlot
            from reportlab.lib.colors import HexColor
            drawing=Drawing(660,250);chart=LinePlot();chart.x=60;chart.y=45;chart.width=580;chart.height=175
            chart.data=[vals];chart.lines[0].strokeColor=HexColor('#246A7A');chart.lines[0].strokeWidth=2
            chart.xValueAxis.valueMin=vals[0][0];chart.xValueAxis.valueMax=vals[-1][0]
            chart.xValueAxis.valueSteps=sorted(set([vals[0][0],*[y for y,v in vals if y%5==0],vals[-1][0]]))
            chart.xValueAxis.labelTextFormat='%d';chart.yValueAxis.labelTextFormat=lambda x:fmt(x)
            chart.yValueAxis.visibleGrid=True;chart.yValueAxis.gridStrokeColor=HexColor('#DDDDDD')
            chart.xValueAxis.labels.fontSize=9;chart.yValueAxis.labels.fontSize=9
            drawing.add(chart);drawing.add(String(60,232,'Habitants',fontSize=10));drawing.add(String(345,10,'Any',fontSize=10))
            png=chart_png(drawing,'population_history');d.add_picture(str(png),width=Cm(16.5))
            para(f"Població total de la mateixa fitxa: {fmt(vals[0][1])} habitants el {vals[0][0]} i {fmt(vals[-1][1])} el {vals[-1][0]}. Gràfic elaborat a partir dels valors publicats; no combina altres padrons.")
    source_line(f)

heading('Medi ambient territori i habitatge',2)
para("El catàleg també incorpora la informació ambiental i territorial disponible. Els espais verds, els usos del sòl i la meteorologia poden orientar on i quan programar activitats; les dades agregades no acrediten l'accessibilitat real d'un itinerari, l'exposició individual a la calor ni la qualitat d'un habitatge. Aquestes condicions s'han de contrastar sobre el terreny i amb les persones participants.")
envrows=[];envsources=[]
for fragment, wanted, unit in [('superficie-despais-verds-en-sol-urba',None,'%'),('superficie-despais-verds-per-habitant',None,'m² per habitant'),('temperatura-maxima-absoluta-txx','Juny','°C'),('precipitacio-acumulada-ppt','Total acumulat','mm')]:
    f=find_fiche(fragment)
    if not f or any(a['id']==f['id'] for a in envsources):continue
    yy,cols,rows=latest_summary(f,30)
    if wanted:rows=[r for r in rows if str(r[0]).strip()==wanted]
    else:rows=rows[:1]
    if rows:
        envrows.append([f['titol']+(f' · {wanted}' if wanted else ''),str(yy or 'Vegeu fitxa'),' / '.join(str(r[-1]).replace('.',',')+(' '+unit if unit not in str(r[-1]) else '') for r in rows)])
        envsources.append(f)
if envrows:table(['Fitxa','Any','Valors publicats'],envrows,[6.5,2,9],8.5)
for f in envsources:source_line(f)
if envsources:para('Meteorologia: precipitació des de l’1/6/2016 (any incomplet); temperatures des de 2017. Canvi d’estació el 2022: de la plaça de la Vila a la plaça Sant Cugat. Cal contextualitzar les comparacions entre anys. Són registres ambientals, no previsions per a una activitat.')
table(['Habitatges a la taula de padró','Recompte','Any'],[['Amb persones empadronades','10.526','2025'],['Sense persones empadronades','1.098','2025'],['Total','11.624','2025']],[10,5,2.5])
para("En aquesta taula, els percentatges publicats no concorden amb els recomptes i el total percentual publicat és 101%. La síntesi reprodueix els recomptes i no reutilitza aquests percentatges. Un habitatge sense persones empadronades no és necessàriament un habitatge buit. Les taules originals es conserven amb l'advertiment corresponent.")
p=para('Font: ');link(p,'Habitatges amb persones empadronades',values['habitatges_padro']['url'])

heading('Propostes per planificar activitats',2)
proposals=[
('Infància i famílies',"Acordar amb pediatria, Salut i Escola, PEE, SOAF i entitats de famílies activitats i horaris adequats. Els 6.455 residents de 0 a 15 anys dimensionen el municipi; les places i la captació s'han d'ajustar a la població realment destinatària.","Assistència, continuïtat i valoració de les famílies."),
('Envelliment i suport social',"Treballar amb Gent Gran, Km 0, la Taula de Soledat, Càritas i farmàcies per explorar preferències, participació i suport. Les 2.133 llars unipersonals són context; no són 2.133 casos de soledat.","Participació sostinguda i suport percebut per les persones participants."),
('Acollida i participació',"Co-dissenyar amb Acollida, mediació i entitats canals comprensibles, horaris i llocs de trobada. Preguntar directament per les barreres i les llengües preferides.","Persones que poden participar, barreres identificades i respostes acordades."),
('Activitat als barris i entorn',"Contrastar la població dels barris amb equipaments, rutes, ombra i accessibilitat. Programar amb agents locals i revisar meteorologia i condicions reals abans de cada activitat.","Assistència per barri, accessibilitat percebuda i incidències."),
]
for title,action,measure in proposals:
    heading(title,3);para(action);para('Seguiment proposat: '+measure)
para("La Comissió de Salut Comunitària hauria de triar les activitats després de contrastar necessitats, capacitat dels actors, participació i viabilitat. El sociograma permet identificar possibles col·laboradors; una proposta de col·laboració no crea una relació documentada.")

# Move the new chapter before the first retained figure, leaving original source intact.
body=d.element.body
chapter=[e for e in list(body)[start-1:] if e.tag!=qn('w:sectPr')]
anchor=original[3]._p
for e in chapter:anchor.addprevious(e)

# Complete catalogue coverage annex with one source link per fiche.
heading('Annex D Catàleg complet de l’Observatori').paragraph_format.page_break_before=True
para(f"Catàleg consultat el 17 d'octubre de 2026: {len(fiches)} fitxes. La cobertura indica les taules efectivament recuperades i els anys que contenen. Una fitxa sense taula queda identificada; no es tracta com un zero ni s'hi imputen dades. Els enllaços del títol obren la fitxa oficial.")
for category in ['Demografia','Medi ambient','Territori']:
    group=[f for f in fiches if f.get('categoria','').casefold()==category.casefold()]
    if not group:group=[f for f in fiches if category.casefold().replace(' ','-') in f.get('categoria','').casefold()]
    if not group:continue
    heading(f'{category} amb {len(group)} fitxes',2)
    rows=[]
    for f in group:
        years=f.get('cobertura',{}).get('anys',[])
        years=sorted(int(y) for y in years if str(y).isdigit())
        period=str(years[0]) if len(years)==1 else f'{years[0]} a {years[-1]}' if years else 'Sense any explícit'
        tables=f.get('cobertura',{}).get('taules',sum(len(v.get('tables',[])) for v in f.get('variants',[])))
        status=f'{tables} taules' if f.get('estat')=='disponible' else 'Sense valors publicats'
        rows.append([f['titol'],period,status])
    t=table(['Fitxa i font oficial','Anys disponibles','Cobertura'],rows,[11.0,3.0,3.5],8.5)
    for row,f in zip(t.rows[1:],group):
        p=row.cells[0].paragraphs[0];p.text='';link(p,f['titol'],f['url'],8.5)
        if f.get('font'):
            p.add_run('\nFont: '+str(f['font'])).font.size=Pt(8)
        if f.get('notes'):
            p.add_run('\nNota: '+' '.join(str(x) for x in f['notes'])).font.size=Pt(8)
    para('Les taules poden incloure variants municipals, territorials, de grup o de temporalitat; el nombre de taules no és un nombre d’indicadors independents.')
heading('Consulta i reutilització',2)
p=para('Visor i descàrregues: ');link(p,'Sociograma de Salt','https://cvragar.github.io/sociograma-salt/#v=observatori')
p=para('Catàleg oficial: ');link(p,'Observatori de Salt','https://observatori.salt.cat/indicadors/')
para("La reproducció de les dades manté l'atribució a l'Ajuntament de Salt i a les fonts indicades a cada fitxa. El portal indica la llicència Creative Commons BY-NC-SA 3.0 ES. Els valors originals, incloses les absències i les incoherències detectades, es conserven als fitxers de dades per facilitar-ne la revisió.")

# Consistent typography; preserve original geometry and illustrations.
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    d.styles[name].font.color.rgb=RGBColor(0,0,0)
    if name=='Title':
        for border in d.styles[name].element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
for border in original[0]._p.xpath('./w:pPr/w:pBdr'):
    border.getparent().remove(border)
for p in d.paragraphs:
    if p.style.name in ['Title','Subtitle'] or p.style.name.startswith('Heading'):
        for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
for t in original_tables:
    if t._tbl.getparent() is not None:format_table(t,size=9 if len(t.columns)<5 else 8.3)
# Page numbering is functional navigation in a lengthy report.
footer=d.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.text='Sociograma Salt · 17/10/2026 · '
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(8)
d.core_properties.title='Sociograma comunitari dels CAP Salt 1 i Salt 2 amb dades de l’Observatori'
d.core_properties.subject='Actualització de dades municipals i catàleg complet'
d.core_properties.author=''
d.core_properties.last_modified_by=''
OUT.parent.mkdir(parents=True,exist_ok=True);d.save(OUT)
assert hashlib.sha256(SRC.read_bytes()).hexdigest()==source_hash
report={'output':str(OUT),'actors':len(actors),'relations':len(rels),'hypotheses':len(hyp),'catalogue_fiches':len(fiches),'with_tables':withdata,'reference_sha256':source_hash,'selected_fiches':list(used)}
(TMP/'build-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
